import pandas as pd # type: ignore
from docx import Document
from docx2pdf import convert
import os
from tkinter import Tk
from tkinter.filedialog import askopenfilename

# Initialize Tkinter root
root = Tk()
root.withdraw()  # Hide the Tkinter GUI window

# Ask the user to select the Excel file
excel_file = askopenfilename(title="Select the Excel file", filetypes=[("Excel files", "*.xlsx")])
if not excel_file:
    print("No Excel file selected. Exiting.")
    exit()

# Ask the user to select the Word template
word_template = askopenfilename(title="Select the Word template", filetypes=[("Word files", "*.docx")])
if not word_template:
    print("No Word template selected. Exiting.")
    exit()

# Load the Excel file
df = pd.read_excel(excel_file)

# Directory to save the generated Word and PDF files
output_dir = 'generated_certificates'
if not os.path.exists(output_dir):
    os.makedirs(output_dir)

for index, row in df.iterrows():
    # Open the Word template
    doc = Document(word_template)
    
    # Replace placeholders in the Word document
    for paragraph in doc.paragraphs:
        if '{name}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{name}', row['Name'])
        if '{department}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{department}', row['Department'])
        if '{roll_no}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{roll_no}', str(row['Roll No']))
    
    # Save the updated document with the roll number as the filename
    docx_filename = f"{output_dir}/{row['Roll No']}.docx"
    doc.save(docx_filename)
    
    # Convert the .docx to .pdf
    pdf_filename = f"{output_dir}/{row['Roll No']}.pdf"
    convert(docx_filename, pdf_filename)
    
    # Optionally, remove the .docx file after conversion to save space
    os.remove(docx_filename)

print("Certificates generated successfully and saved as PDFs.")
